#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""每个工作日 5 点管线选股 → Excel → 企业微信推送。

数据源: output/daily_recommend.json（5 点管线产物，06:14 左右完成）
  - 注意: 09:35 会被 live_momentum_scanner 重排覆盖为 Top 版，
    因此本脚本 cron 必须定在 06:20（完整版，含 full_candidate_pool）。
  - 若发现 recommendations 太少（疑似被覆盖），回退到 full_candidate_pool。

产出:
  output/excel_picks/AlphaPilot选股_YYYY-MM-DD.xlsx  — Excel（Sheet1 全部 / Sheet2 Top10 / Sheet3 Top2+说明）
  企业微信群推送: markdown 摘要 + Excel 文件

幂等: 同一天只发一次（标记文件 output/excel_picks/.sent_YYYY-MM-DD）。
  --force 可强制重发。

建议 cron: 工作日 06:20
  20 6 * * 1-5 cd /home/ubuntu/alphapilot && python3 -u scripts/send_daily_picks_excel.py >> output/logs/send_daily_picks_excel.log 2>&1
"""
from __future__ import annotations

import argparse
import json
import os
import sys
from datetime import datetime
from pathlib import Path

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

ROOT = Path(os.environ.get("ALPHAPILOT_ROOT") or "/home/ubuntu/alphapilot")
if not (ROOT / "output" / "daily_recommend.json").exists():
    # 本地开发或路径缺失时回退到脚本目录
    candidate = Path(__file__).resolve().parents[1]
    if (candidate / "output" / "daily_recommend.json").exists():
        ROOT = candidate
os.chdir(ROOT)
sys.path.insert(0, str(ROOT))
sys.path.insert(0, str(ROOT / "scripts"))

REC = ROOT / "output" / "daily_recommend.json"
MORNING = ROOT / "output" / "morning_live_picks.json"
OUT_DIR = ROOT / "output" / "excel_picks"


def log(msg: str) -> None:
    print(f"[{datetime.now().strftime('%Y-%m-%d %H:%M:%S')}] {msg}", flush=True)


def _fmt(v, nd=2):
    if v is None:
        return None
    try:
        f = float(v)
        return round(f, nd)
    except (TypeError, ValueError):
        return v


def _fmt_wan(v):
    """主力净额 → 万元，带符号。"""
    if v is None:
        return None
    try:
        f = float(v)
        return round(f / 10000.0, 1)
    except (TypeError, ValueError):
        return None


def load_records() -> list[dict]:
    """读取今日选股。优先 recommendations；若太少则回退 full_candidate_pool。"""
    if not REC.exists():
        return []
    d = json.loads(REC.read_text(encoding="utf-8"))
    recs = d.get("recommendations") or []
    if len(recs) < 30:
        full = d.get("full_candidate_pool") or []
        if len(full) > len(recs):
            log(f"recommendations 疑似被覆盖({len(recs)}只)，回退 full_candidate_pool({len(full)}只)")
            recs = full
    return recs


def load_top2_from_daily(d: dict) -> tuple[list[dict], str]:
    """从 daily_recommend.json 提取下单 Top2（morning_pick_rank 非空）。

    返回 (top2, top2_date)。morning_pick_rank 由 09:35 重排写回，日期为今日。
    """
    recs = d.get("recommendations") or []
    ranked = [r for r in recs if r.get("morning_pick_rank") is not None]
    ranked.sort(key=lambda r: r["morning_pick_rank"])
    return ranked[:2], datetime.now().strftime("%Y-%m-%d")


def load_top2_from_morning() -> tuple[list[dict], str]:
    """回退：从 morning_live_picks.json 提取 Top2。

    返回 (top2, top2_date)，top2_date 取自该文件的 asof 日期
    （可能是今日 09:35 已更新，也可能是上一交易日未更新的旧终选）。
    """
    if not MORNING.exists():
        return [], ""
    try:
        m = json.loads(MORNING.read_text(encoding="utf-8"))
    except Exception:
        return [], ""
    picks = m.get("picks") or m.get("recommendations") or []
    if not isinstance(picks, list):
        return [], ""
    picks = [p for p in picks if isinstance(p, dict)]
    asof = m.get("asof") or ""
    top2_date = asof[:10] if len(asof) >= 10 else ""
    return picks[:2], top2_date


def make_workbook(records: list[dict], top2: list[dict], meta: dict, today: str, top2_date: str = "") -> Path:
    """生成 Excel 文件，返回路径。top2_date 为下单 Top2 实际终选日期。"""
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    path = OUT_DIR / f"AlphaPilot选股_{today}.xlsx"

    wb = openpyxl.Workbook()

    header_fill = PatternFill("solid", fgColor="1F4E79")
    header_font = Font(color="FFFFFF", bold=True, size=10)
    thin = Side(style="thin", color="CCCCCC")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    top2_fill = PatternFill("solid", fgColor="FFF2CC")
    title_font = Font(bold=True, size=13, color="1F4E79")

    def style_header(ws, headers):
        for c, h in enumerate(headers, 1):
            cell = ws.cell(row=1, column=c, value=h)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            cell.border = border
        ws.freeze_panes = "A2"

    def auto_width(ws, min_w=8, max_w=28):
        for col in ws.columns:
            letter = get_column_letter(col[0].column)
            width = max(min_w, min(max_w, max(len(str(c.value or "")) for c in col) + 2))
            ws.column_dimensions[letter].width = width

    # ---------- Sheet1 全部选股 ----------
    ws1 = wb.active
    ws1.title = "全部选股"
    headers1 = [
        "排名", "代码", "名称", "综合评分", "原始模型分", "现价", "建议买入价", "目标价", "止损价",
        "涨跌幅%", "主力净额(万)", "主动买占比", "资金阶段", "换手率%", "量比", "所属板块",
        "选股臂", "板块热度", "盘前竞价", "盘前跳空%", "LLM情绪", "备注(LLM理由)",
    ]
    ws1.append(headers1)
    style_header(ws1, headers1)

    sorted_recs = sorted(records, key=lambda r: -(r.get("score") if r.get("score") is not None else r.get("score_raw") or 0))
    top2_syms = {t.get("symbol") for t in top2}

    for i, r in enumerate(sorted_recs, 1):
        reason = (r.get("llm_reason") or "")[:120]
        row = [
            i,
            r.get("symbol"),
            r.get("name"),
            _fmt(r.get("score")),
            _fmt(r.get("score_raw") if r.get("score_raw") is not None else r.get("ml_score")),
            _fmt(r.get("price")),
            _fmt(r.get("buy_price")),
            _fmt(r.get("target_price")),
            _fmt(r.get("stop_price")),
            _fmt(r.get("change_pct")),
            _fmt_wan(r.get("main_net")),
            _fmt(r.get("active_buy_ratio"), 3),
            r.get("money_phase_label"),
            _fmt(r.get("turnover")),
            _fmt(r.get("volume_ratio")),
            r.get("industry_l1"),
            r.get("selection_arm") or r.get("arm"),
            _fmt(r.get("sector_heat"), 3),
            r.get("pre_market_action"),
            _fmt(r.get("pre_market_gap_pct")),
            r.get("llm_sentiment"),
            reason,
        ]
        ws1.append(row)
        if r.get("symbol") in top2_syms:
            for c in range(1, len(headers1) + 1):
                ws1.cell(row=i + 1, column=c).fill = top2_fill
        for c in range(1, len(headers1) + 1):
            ws1.cell(row=i + 1, column=c).border = border
    auto_width(ws1)

    # ---------- Sheet2 Top10 精选 ----------
    ws2 = wb.create_sheet("Top10精选")
    headers2 = ["排名", "代码", "名称", "综合评分", "原始模型分", "现价", "建议买入价", "目标价", "止损价",
                "主力净额(万)", "主动买占比", "资金阶段", "所属板块", "选股臂", "备注(LLM理由)"]
    ws2.append(headers2)
    style_header(ws2, headers2)
    for i, r in enumerate(sorted_recs[:10], 1):
        ws2.append([
            i, r.get("symbol"), r.get("name"),
            _fmt(r.get("score")), _fmt(r.get("score_raw") if r.get("score_raw") is not None else r.get("ml_score")),
            _fmt(r.get("price")), _fmt(r.get("buy_price")), _fmt(r.get("target_price")), _fmt(r.get("stop_price")),
            _fmt_wan(r.get("main_net")), _fmt(r.get("active_buy_ratio"), 3), r.get("money_phase_label"),
            r.get("industry_l1"), r.get("selection_arm") or r.get("arm"), (r.get("llm_reason") or "")[:100],
        ])
        if r.get("symbol") in top2_syms:
            for c in range(1, len(headers2) + 1):
                ws2.cell(row=i + 1, column=c).fill = top2_fill
        for c in range(1, len(headers2) + 1):
            ws2.cell(row=i + 1, column=c).border = border
    auto_width(ws2)

    # ---------- Sheet3 Top2下单 + 说明 ----------
    ws3 = wb.create_sheet("Top2下单说明")
    ws3.cell(row=1, column=1, value=f"AlphaPilot 当日选股说明  {today}").font = title_font

    def kv(r, k, v):
        ws3.cell(row=r, column=1, value=k).font = Font(bold=True)
        ws3.cell(row=r, column=2, value=v)

    r = 3
    kv(r, "选股日期", today); r += 1
    kv(r, "管线运行时间", meta.get("run_at")); r += 1
    kv(r, "协议", meta.get("protocol") or meta.get("scanner")); r += 1
    kv(r, "模型", meta.get("model")); r += 1
    kv(r, "特征维度", meta.get("n_features")); r += 1
    kv(r, "全部选股数", len(sorted_recs)); r += 1
    kv(r, "建议下单数(TopN)", meta.get("recommend_top_n")); r += 1
    kv(r, "推荐池大小", meta.get("recommend_pool_n")); r += 1
    kv(r, "市场环境", json.dumps(meta.get("market_env_flags"), ensure_ascii=False)); r += 1
    r += 1
    if top2_date and top2_date != today:
        ws3.cell(row=r, column=1,
                 value=f"下单 Top2 · {top2_date} 终选（⚠️ 今日 09:35 终选尚未更新，下表为上一交易日终选）").font = Font(bold=True, size=11)
    else:
        ws3.cell(row=r, column=1, value=f"当日下单 Top2 · {top2_date or today} 终选").font = Font(bold=True, size=11)
    r += 1
    ws3.cell(row=r, column=1, value="排名").font = Font(bold=True)
    ws3.cell(row=r, column=2, value="代码").font = Font(bold=True)
    ws3.cell(row=r, column=3, value="名称").font = Font(bold=True)
    ws3.cell(row=r, column=4, value="综合评分").font = Font(bold=True)
    ws3.cell(row=r, column=5, value="建议买入价").font = Font(bold=True)
    ws3.cell(row=r, column=6, value="现价").font = Font(bold=True)
    r += 1
    for idx, t in enumerate(top2, 1):
        price_txt = t.get("price")
        if price_txt is None:
            price_txt = t.get("buy_price")
        kv_row = [
            idx, t.get("symbol"), t.get("name"),
            _fmt(t.get("score")), _fmt(t.get("buy_price")), _fmt(price_txt),
        ]
        for c, val in enumerate(kv_row, 1):
            ws3.cell(row=r, column=c, value=val)
        r += 1
    r += 1
    note = (
        "口径说明：\n"
        "1. 综合评分 = 模型分 + 资金流 + 板块热度等综合（与『评分 Top10 · 09:35 定格』同口径）。\n"
        "2. 全部选股来自 5 点管线全链（含模型评分/资金门/竞价门/板块研究等层层筛选）。\n"
        "3. 现价为最近收盘价；实际买入按盘中 5 分钟线 VWAP 回踩低位执行。\n"
        "4. 此表为量化研究参考，不构成投资建议。"
    )
    for line in note.splitlines():
        ws3.cell(row=r, column=1, value=line)
        r += 1
    ws3.column_dimensions["A"].width = 22
    ws3.column_dimensions["B"].width = 30
    ws3.column_dimensions["C"].width = 14
    ws3.column_dimensions["D"].width = 12
    ws3.column_dimensions["E"].width = 14
    ws3.column_dimensions["F"].width = 12

    wb.save(path)
    log(f"Excel 已生成: {path}")
    return path


def build_summary_md(records: list[dict], top2: list[dict], today: str, meta: dict, top2_date: str = "") -> str:
    """生成 markdown 摘要（推送用）。"""
    sorted_recs = sorted(records, key=lambda r: -(r.get("score") if r.get("score") is not None else r.get("score_raw") or 0))
    lines = [
        f"## 📈 AlphaPilot 今日选股 · {today}",
        "",
        f"> 5 点管线选出 **{len(sorted_recs)}** 只，模型 {meta.get('model')}（{meta.get('n_features')} 维）",
        "",
    ]
    if top2:
        if top2_date and top2_date != today:
            lines.append(f"**🎯 下单 Top2 · {top2_date} 终选**")
            lines.append(f"> ⚠️ 今日 09:35 终选尚未更新，以下为 **{top2_date}**（上一交易日）终选结果")
        else:
            lines.append(f"**🎯 当日下单 Top2 · {today} 终选**")
        for idx, t in enumerate(top2, 1):
            price_txt = t.get("price")
            if price_txt is None:
                price_txt = t.get("buy_price")
            lines.append(
                f"{idx}. **{t.get('name')}**（{t.get('symbol')}） 评分 {t.get('score')} | "
                f"建议买入 {t.get('buy_price')} | 现价 {price_txt}"
            )
        lines.append("")
    lines.append(f"**🏆 综合评分 Top5 · {today}**")
    for i, r in enumerate(sorted_recs[:5], 1):
        lines.append(
            f"{i}. {r.get('name')}（{r.get('symbol')}） {_fmt(r.get('score'))} | "
            f"{r.get('industry_l1')} | {r.get('money_phase_label') or ''}"
        )
    lines.append("")
    lines.append(f"📎 完整清单见 Excel（共 {len(sorted_recs)} 只，含全部评分/资金/板块字段）")
    return "\n".join(lines)


if __name__ == "__main__":
    ap = argparse.ArgumentParser()
    ap.add_argument("--force", action="store_true", help="同一天强制重发")
    ap.add_argument("--no-send", action="store_true", help="只生成 Excel，不推送")
    args = ap.parse_args()

    today = datetime.now().strftime("%Y-%m-%d")
    marker = OUT_DIR / f".sent_{today}"
    if marker.exists() and not args.force:
        log(f"今日({today})已发送过，跳过（--force 可重发）")
        sys.exit(0)

    if not REC.exists():
        log("daily_recommend.json 不存在，跳过")
        sys.exit(0)

    data = json.loads(REC.read_text(encoding="utf-8"))
    meta = {
        "run_at": data.get("run_at"),
        "protocol": data.get("protocol"),
        "scanner": data.get("scanner"),
        "model": None,
        "n_features": None,
        "recommend_top_n": data.get("recommend_top_n"),
        "recommend_pool_n": data.get("recommend_pool_n"),
        "market_env_flags": data.get("market_env_flags"),
    }
    recs = data.get("recommendations") or []
    if recs and isinstance(recs[0], dict):
        meta["model"] = recs[0].get("model")
        meta["n_features"] = recs[0].get("n_features")

    records = load_records()
    if not records:
        log("今日无选股记录，跳过")
        sys.exit(0)

    top2, top2_date = load_top2_from_daily(data)
    if not top2:
        top2, top2_date = load_top2_from_morning()
    if not top2_date:
        top2_date = today

    xlsx = make_workbook(records, top2, meta, today, top2_date)

    if args.no_send:
        log("--no-send，跳过推送")
        sys.exit(0)

    from wecom_push import send_markdown, send_file

    md = build_summary_md(records, top2, today, meta, top2_date)
    ok1, err1 = send_markdown(md)
    log(f"发送 markdown 摘要 -> ok={ok1} err={err1}")
    ok2, err2 = send_file(xlsx)
    log(f"发送 Excel 文件 -> ok={ok2} err={err2}")

    if ok1 or ok2:
        marker.write_text(datetime.now().isoformat(), encoding="utf-8")
        log("已记录发送标记")
        sys.exit(0)
    log("推送失败，不记录标记（下次 cron 会重试）")
    sys.exit(1)
